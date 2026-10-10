import io
import sys
from types import ModuleType
from datetime import date

import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest

from app_modules.modules.reportes_page import (
    _excel_bytes,
    _printable_html,
    _report_subtitle,
    _safe_filename,
)


SCRIPT = '''
import streamlit as st
from app_modules.modules.reportes_page import render
from lotes_test_api import api_get, api_post, api_put, api_delete, ROLE

def flash_set(tab, kind, msg): ...
def flash_show(tab): ...

st.session_state.auth = {"rol": ROLE}
render(
    st=st, tabs=None, selected_tab="reportes", rol=ROLE, API="http://mock-api",
    auth_headers=None, api_get=api_get, api_post=api_post, api_put=api_put,
    api_delete=api_delete, flash_show=flash_show, flash_set=flash_set,
    get_jwt=lambda: "test-jwt", show_printers_panel=None,
    COLUMNAS_LISTADO=None, COLUMNAS_IMPRESION=None,
)
'''


# -------------------------
# Funciones puras
# -------------------------
def _sample_df():
    return pd.DataFrame([
        {"dni": "11111111", "persona": "PEREZ RUIZ ANA", "empacador": 12, "seleccionador": 0, "total": 12},
        {"dni": "22222222", "persona": "SOTO LUIS", "empacador": 0, "seleccionador": 8, "total": 8},
    ])


def test_excel_bytes_generates_valid_xlsx_with_header_and_rows():
    sub = _report_subtitle("Lote: LOTE-1", "", "PALTA")
    data = _excel_bytes(sub, _sample_df())
    assert data[:2] == b"PK"  # firma ZIP/XLSX
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(data))
    sheet = book["Reporte"]
    assert sheet["A1"].value == "Sistema de Etiquetas QR — Reporte de producción"
    assert sheet["A2"].value == "Lote: LOTE-1"
    assert sheet["A3"].value == "Producto: PALTA"
    # La tabla arranca tras el bloque de detalle: encabezados y filas en su lugar.
    # (startrow es 0-based: con 4 líneas de detalle, la tabla empieza en la fila 7.)
    assert sheet["A7"].value == "dni"
    assert sheet["A8"].value == "11111111"


def test_excel_and_print_accept_date_range_filters():
    filtro = "Fechas: 02/10/2026 al 09/10/2026"
    sub = _report_subtitle(filtro, "", "")
    data = _excel_bytes(sub, _sample_df())
    from openpyxl import load_workbook
    sheet = load_workbook(io.BytesIO(data))["Reporte"]
    assert sheet["A2"].value == filtro

    html = _printable_html("Reporte de producción por DNI", sub,
                           [("Total lecturas", 20)], _sample_df())
    assert "Fechas: 02/10/2026 al 09/10/2026" in html
    assert "PEREZ RUIZ ANA" in html


def test_printable_html_contains_title_totals_and_table():
    sub = _report_subtitle("Lote: LOTE-1", "", "")
    html = _printable_html("Reporte de producción por DNI", sub,
                           [("Total lecturas", 20)], _sample_df())
    assert html.startswith("<!DOCTYPE html>")
    assert "Reporte de producción por DNI" in html
    assert "LOTE-1" in html
    assert "PEREZ RUIZ ANA" in html
    assert "11111111" in html
    assert "<table" in html and "<th>" in html
    # No debe haber etiquetas rotas: el documento cierra correctamente.
    assert html.rstrip().endswith("</html>")


def test_safe_filename_sanitizes_filters():
    assert "/" not in _safe_filename("reporte_dni", "Lote: LO/TE 1-2")
    assert "\\" not in _safe_filename("reporte_dni", "Fechas: \\x")
    assert _safe_filename("reporte_dni", "LOTE-9").startswith("reporte_dni_LOTE-9")
    assert _safe_filename("reporte_dni", "LOTE-9").endswith(".xlsx")


# -------------------------
# Página completa (AppTest)
# -------------------------
@pytest.fixture
def ui_api(client, user, monkeypatch):
    module = ModuleType("lotes_test_api")
    module.ROLE = user["rol"]
    module.api_get = lambda path, **kwargs: client.get("/api" + path, **kwargs)
    module.api_post = lambda path, **kwargs: client.post("/api" + path, **kwargs)
    module.api_put = lambda path, **kwargs: client.put("/api" + path, **kwargs)
    module.api_delete = lambda path: client.delete("/api" + path)
    monkeypatch.setitem(sys.modules, "lotes_test_api", module)
    return module


class _FakeResponse:
    def __init__(self, payload):
        self.status_code = 200
        self._payload = payload

    def json(self):
        return self._payload


def _patch_reports(monkeypatch, payload):
    import app_modules.modules.reportes_page as page
    captured = []

    def fake_get(url, params=None, headers=None, timeout=None):
        captured.append({"url": url, "params": dict(params or {})})
        return _FakeResponse(payload)

    monkeypatch.setattr(page.requests, "get", fake_get)
    return captured


def test_consultar_by_lote_sends_lote_param_and_persists(ui_api, client, monkeypatch):
    client.post("/api/lotes", json={"codigo": "REP-1"})
    payload = {
        "producto": None, "lote_codigo": "REP-1",
        "totals": {"total_lecturas": 20, "emp_lecturas": 12, "sel_lecturas": 8},
        "rows": [{"dni": "11111111", "persona": "PEREZ RUIZ ANA",
                  "empacador": 12, "seleccionador": 0, "total": 12}],
    }
    captured = _patch_reports(monkeypatch, payload)

    at = AppTest.from_string(SCRIPT).run(timeout=20)
    assert not at.exception
    # Modo por defecto: Lote.
    assert at.radio[0].value == "Lote"
    next(b for b in at.button if b.label == "🔍 Consultar").click()
    at = at.run(timeout=20)
    assert not at.exception

    assert captured[-1]["params"]["lote_codigo"] == "REP-1"
    assert at.session_state["rep_dni_data"]["totals"]["total_lecturas"] == 20
    assert at.session_state["rep_dni_filtro"] == "Lote: REP-1"
    assert any("Reporte por DNI" in m.value for m in at.markdown)
    assert len(at.dataframe) >= 1

    # Una segunda re-ejecución (p. ej. al descargar) conserva el reporte.
    at = at.run(timeout=20)
    assert at.session_state["rep_dni_data"] is not None


def test_consultar_by_date_range_sends_dates_without_lote(ui_api, client, monkeypatch):
    client.post("/api/lotes", json={"codigo": "REP-F"})
    payload = {
        "producto": None, "lote_codigo": None,
        "date_from": "2026-10-02", "date_to": "2026-10-09",
        "totals": {"total_lecturas": 5, "emp_lecturas": 3, "sel_lecturas": 2},
        "rows": [{"dni": "33333333", "persona": "QUISPE MARIA",
                  "empacador": 3, "seleccionador": 2, "total": 5}],
    }
    captured = _patch_reports(monkeypatch, payload)

    at = AppTest.from_string(SCRIPT).run(timeout=20)
    assert not at.exception
    # Cambiar a modo fechas y fijar el rango.
    at.radio[0].set_value("Rango de fechas")
    at = at.run(timeout=20)
    assert not at.exception
    at.date_input(key="rep_desde").set_value(date(2026, 10, 2))
    at.date_input(key="rep_hasta").set_value(date(2026, 10, 9))
    at = at.run(timeout=20)
    next(b for b in at.button if b.label == "🔍 Consultar").click()
    at = at.run(timeout=20)
    assert not at.exception

    sent = captured[-1]["params"]
    assert sent["date_from"] == "2026-10-02"
    assert sent["date_to"] == "2026-10-09"
    assert "lote_codigo" not in sent
    assert at.session_state["rep_dni_filtro"].startswith("Fechas:")
    assert "02/10/2026" in at.session_state["rep_dni_filtro"]


def test_operators_report_by_date_range(ui_api, client, monkeypatch):
    client.post("/api/lotes", json={"codigo": "REP-O"})
    payload = {
        "producto": None, "lote_codigo": None, "rows": [
            {"user_id": "op1", "total": 15, "dnis_distintos": 3, "ultima_lectura": "2026-01-05"},
        ],
    }
    captured = _patch_reports(monkeypatch, payload)

    at = AppTest.from_string(SCRIPT).run(timeout=20)
    at.radio[0].set_value("Rango de fechas")
    at = at.run(timeout=20)
    next(b for b in at.button if b.label == "📋 Reporte de operadores").click()
    at = at.run(timeout=20)
    assert not at.exception
    assert captured[-1]["params"]["date_from"]
    assert at.session_state["rep_op_data"]["rows"][0]["user_id"] == "op1"
    assert any("Reporte por operadores" in m.value for m in at.markdown)


def test_date_range_validation_blocks_inverted_dates(ui_api, client, monkeypatch):
    _patch_reports(monkeypatch, {"rows": []})
    at = AppTest.from_string(SCRIPT).run(timeout=20)
    at.radio[0].set_value("Rango de fechas")
    at = at.run(timeout=20)
    at.date_input(key="rep_desde").set_value(date(2026, 10, 9))
    at.date_input(key="rep_hasta").set_value(date(2026, 10, 2))
    at = at.run(timeout=20)
    assert not at.exception
    consultar = next(b for b in at.button if b.label == "🔍 Consultar")
    assert consultar.disabled
    assert any("fecha final" in w.value.lower() for w in at.warning)


def test_missing_openpyxl_shows_warning_instead_of_crashing(ui_api, client, monkeypatch):
    import app_modules.modules.reportes_page as page
    monkeypatch.setattr(page, "_EXCEL_DISPONIBLE", False)
    client.post("/api/lotes", json={"codigo": "REP-M"})
    payload = {
        "producto": None, "lote_codigo": "REP-M",
        "totals": {"total_lecturas": 9, "emp_lecturas": 5, "sel_lecturas": 4},
        "rows": [{"dni": "11111111", "persona": "PEREZ RUIZ ANA",
                  "empacador": 5, "seleccionador": 0, "total": 5}],
    }
    _patch_reports(monkeypatch, payload)

    at = AppTest.from_string(SCRIPT).run(timeout=20)
    assert not at.exception
    next(b for b in at.button if b.label == "🔍 Consultar").click()
    at = at.run(timeout=20)
    assert not at.exception  # sin openpyxl la página no se rompe
    assert any("reporte por DNI" in m.value.lower() or "Reporte por DNI" in m.value
               for m in at.markdown)
    assert any("openpyxl" in w.value for w in at.warning)


def test_report_buttons_stay_disabled_without_lotes_in_lote_mode(ui_api, monkeypatch):
    _patch_reports(monkeypatch, {"rows": []})
    at = AppTest.from_string(SCRIPT).run(timeout=20)
    assert not at.exception
    assert at.radio[0].value == "Lote"
    consultar = next(b for b in at.button if b.label == "🔍 Consultar")
    assert consultar.disabled
    # En modo fechas el botón se habilita aunque no existan lotes.
    at.radio[0].set_value("Rango de fechas")
    at = at.run(timeout=20)
    assert not next(b for b in at.button if b.label == "🔍 Consultar").disabled
